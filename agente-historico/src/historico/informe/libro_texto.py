"""Las hojas de texto del libro: Resumen (primera) y Método (ultima)."""
from __future__ import annotations

from openpyxl import Workbook

from historico.informe.libro_estilo import (ANCHO_TEXTO, AJUSTE, TITULO, cabecera,
                                            seccion)
from historico.informe.tabla import Hecho, Hoja, texto
from historico.informe.verificar import HIPOTESIS


def resumen(wb: Workbook, meta: dict, lectura: dict, advertencias: list[str],
             hechos: list[Hecho]) -> None:
    ws = wb.active
    ws.title = "Resumen"
    ws["A1"], ws["A1"].font = "Informe AgroVoltaic · San Carlos", TITULO
    ws["A2"] = f"Periodo: {meta['desde']} a {meta['hasta']} (ambos incluidos)"
    ws["A3"] = (f"Generado: {meta['generado']} · último dato eléctrico de la base: "
                f"{meta['ultimo_dato'] or 'sin dato'}")
    if meta.get("foco"):
        ws["A4"] = f"Foco pedido: {meta['foco']}"

    fila = seccion(ws, 6, "Lectura")
    if lectura["parrafos"]:
        cabecera(ws, fila, ["tipo", "texto", "hechos_citados"])
        for p in lectura["parrafos"]:
            fila += 1
            ws.cell(row=fila, column=1,
                    value="Hipótesis" if p["tipo"] == HIPOTESIS else "Hecho")
            ws.cell(row=fila, column=2, value=p["texto"]).alignment = AJUSTE
            ws.cell(row=fila, column=3, value=", ".join(p["citas"]))
    else:
        ws.cell(row=fila, column=2,
                value=f"Sin lectura: {lectura['motivo']}").alignment = AJUSTE

    fila = seccion(ws, fila + 2, "Advertencias")
    for a in advertencias:
        ws.cell(row=fila, column=2, value=a).alignment = AJUSTE
        fila += 1

    fila = seccion(ws, fila + 1, "Hechos")
    cabecera(ws, fila, ["id", "indicador", "valor", "origen"])
    for h in hechos:
        fila += 1
        for i, v in enumerate((h.id, h.indicador, texto(h), h.origen), start=1):
            ws.cell(row=fila, column=i, value=v)
    for letra, ancho in (("A", 14), ("B", ANCHO_TEXTO), ("C", 24), ("D", 44)):
        ws.column_dimensions[letra].width = ancho


def metodo(wb: Workbook, hojas: list[Hoja], parametros: list[tuple[str, str]],
            traza: list[tuple[str, str]]) -> None:
    ws = wb.create_sheet("Método")
    fila = seccion(ws, 1, "Parámetros y definiciones")
    for clave, valor in parametros:
        ws.cell(row=fila, column=1, value=clave)
        ws.cell(row=fila, column=2, value=valor).alignment = AJUSTE
        fila += 1
    fila = seccion(ws, fila + 1, "Trazabilidad")
    for clave, valor in traza:
        ws.cell(row=fila, column=1, value=clave)
        ws.cell(row=fila, column=2, value=valor).alignment = AJUSTE
        fila += 1
    fila = seccion(ws, fila + 1, "Notas por hoja")
    for h in hojas:
        ws.cell(row=fila, column=1, value=h.nombre)
        ws.cell(row=fila, column=2, value=f"{h.titulo}. {h.nota}").alignment = AJUSTE
        fila += 1
    fila = seccion(ws, fila + 1, "Diccionario de columnas")
    cabecera(ws, fila, ["hoja", "definicion", "columna", "etiqueta", "unidad"])
    for h in hojas:
        for c in h.columnas:
            fila += 1
            for i, v in enumerate((h.nombre, c.descripcion, c.clave, c.etiqueta,
                                   c.unidad), start=1):
                ws.cell(row=fila, column=i, value=v).alignment = AJUSTE
    for letra, ancho in (("A", 30), ("B", ANCHO_TEXTO), ("C", 36), ("D", 34), ("E", 10)):
        ws.column_dimensions[letra].width = ancho
