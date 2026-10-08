"""
Tests del informe en Excel: sin base de datos y sin LLM.

Lo que protegen, en orden de importancia:

  1. Que la lectura NO pueda publicar una cifra que el libro no tiene.
  2. Que un dia sin dato salga vacio y con motivo, nunca en cero.
  3. Que el libro salga igual cuando el modelo falla o no pasa la revision.
  4. Que las hojas de datos se lean con la cabecera en la fila 1.

Estructura Given-When-Then.
"""
import io
import json
from types import SimpleNamespace

import openpyxl
import pytest
from fastapi.testclient import TestClient

from historico import api, informe
from historico.analitica import rendimiento, resultado
from historico.api import app
from historico.informe import flujo, hechos, redaccion, secciones, verificar
from historico.informe.tabla import Hecho, texto
from historico.tools import preparar_informe

CLIENTE = TestClient(app)


# ── Insumos de prueba: tres dias de calendario ──────────────────────────────────
def _dia_calendario(fecha, filas=150, sin_acoplar=0, bajo_sol=False, veredicto="ok",
                    clase="parcial"):
    return {"fecha": fecha, "filas_electrico": filas, "filas_radiacion": filas,
            "lecturas_sin_acoplar": sin_acoplar, "parada_bajo_sol": bajo_sol,
            "veredicto": veredicto, "clase": clase, "kt_medio": 0.5}


def _fila_pr(dia, e1=5.0, e2=3.0, ghi=4000.0, horas_rad=12.0):
    return {"dia": dia, "horas_sol": 12.0, "horas_rad": horas_rad, "horas_ele": 12.0,
            "ghi_wh_m2": ghi, "poa1_bif_wh_m2": 4400.0, "poa2_bif_wh_m2": 3000.0,
            "poa1_front_wh_m2": 4000.0, "poa2_front_wh_m2": 1500.0,
            "e1_integral_wh": 4300.0, "e2_integral_wh": 2600.0,
            "e1_contador_kwh": e1, "e2_contador_kwh": e2}


def _fila_energia(dia, cierre=7.5):
    return {"dia": dia, "ac_cierre": cierre, "n_ac": 140, "filas": 150}


@pytest.fixture
def insumos():
    """Un dia bueno, uno con la planta parada y media radiacion, y uno sin datos."""
    return {
        "calendario": [
            _dia_calendario("2026-07-01"),
            _dia_calendario("2026-07-02", sin_acoplar=24, bajo_sol=True,
                            veredicto="grave"),
            _dia_calendario("2026-07-03", filas=0, veredicto="sin_datos", clase=None),
        ],
        "filas_pr": [_fila_pr("2026-07-01"),
                     _fila_pr("2026-07-02", e1=0.0, e2=0.0, horas_rad=5.0)],
        "filas_energia": [_fila_energia("2026-07-01"), _fila_energia("2026-07-02", 0.0)],
        "tipos": [{"fuente": "monitoreo_sc_electrico", "tipo": "fuera_de_rango",
                   "severidad": "grave", "dias": 1, "variables": 2, "lecturas": 30,
                   "primer_dia": "2026-07-02", "ultimo_dia": "2026-07-02"}],
        "motivo_poa": resultado.SIN_LECTURAS,
    }


@pytest.fixture
def composicion(insumos):
    return flujo.componer(insumos, "2026-07-01", "2026-07-03", "2026-07-02T17:00:00")


def _dias(insumos):
    return secciones.filas_diario(insumos["calendario"], insumos["filas_pr"],
                                  insumos["filas_energia"])


# ── Diario ──────────────────────────────────────────────────────────────────────
def test_el_diario_tiene_un_renglon_por_dia_de_calendario(insumos):
    # Given/When
    dias = _dias(insumos)

    # Then: el dia sin datos tambien esta, vacio y con su motivo
    assert [d["fecha"] for d in dias] == ["2026-07-01", "2026-07-02", "2026-07-03"]
    vacio = dias[2]
    assert vacio["energia_ac_kwh"] is None and vacio["pr_inclinado_ghi"] is None
    assert vacio["motivo"] == secciones.SIN_FILAS


def test_el_pr_diario_es_el_de_rendimiento(insumos):
    # Given/When
    bueno = _dias(insumos)[0]

    # Then: (5 kWh / 1,42 kWp) / (4 kWh/m2), calculado por el modulo de siempre
    esperado = round((5.0 / rendimiento.KWP_POR_ARREGLO) / 4.0, 3)
    assert bueno["pr_inclinado_ghi"] == esperado
    assert bueno["apto_pr"] == "si" and bueno["motivo"] == ""


def test_un_dia_no_apto_no_lleva_pr_aunque_tenga_energia(insumos):
    # Given: el segundo dia tiene 5 de 12 horas de radiacion
    # When
    malo = _dias(insumos)[1]

    # Then: sin PR, con el motivo en palabras, y la energia en cero REAL (planta parada)
    assert malo["apto_pr"] == "no" and malo["pr_inclinado_ghi"] is None
    assert "cobertura" in malo["motivo"]
    assert malo["energia_ac_kwh"] == 0.0
    assert malo["planta_parada"] == "si" and malo["horas_parada"] == pytest.approx(2.0)


def test_un_dia_apto_con_la_planta_parada_lo_dice_en_el_motivo(insumos):
    # Given: el dia bueno, pero con lecturas sin acoplar
    insumos["calendario"][0]["lecturas_sin_acoplar"] = 12

    # When/Then: su PR es real pero mide la parada
    assert _dias(insumos)[0]["motivo"] == secciones.PARADA


# ── Mensual y hechos ────────────────────────────────────────────────────────────
def test_el_total_suma_solo_lo_que_existe(composicion):
    # Given/When
    total = composicion["hojas"][1].filas[-1]

    # Then: tres dias de calendario, dos con datos, uno apto
    assert total["mes"] == "Total"
    assert (total["dias_calendario"], total["dias_con_datos"],
            total["dias_aptos_pr"]) == (3, 2, 1)
    assert total["energia_ac_kwh"] == pytest.approx(7.5)
    assert total["dias_parada"] == 1


def test_los_hechos_salen_numerados_y_con_origen(composicion):
    # Given/When
    lista = composicion["hechos"]

    # Then
    assert [h.id for h in lista] == [f"H{n}" for n in range(1, len(lista) + 1)]
    por_nombre = {h.indicador: h for h in lista}
    assert texto(por_nombre["Días con datos eléctricos"]) == "2 días"
    assert texto(por_nombre["Días aptos para PR"]) == "1 día"
    assert por_nombre["Energía AC registrada"].origen == "Mensual: fila Total"


def test_con_un_solo_mes_no_se_repiten_hechos_mensuales(composicion):
    assert not any("2026-07" in h.indicador for h in composicion["hechos"])


def test_un_hecho_sin_valor_se_lee_como_sin_dato():
    assert texto(Hecho("H1", "PR", None, "", "", 3)) == "sin dato"
    assert texto(Hecho("H1", "Energía", 1644.02, "kWh", "", 2)) == "1.644,02 kWh"


def test_los_hallazgos_graves_van_antes_que_los_frecuentes():
    # Given: un `info` que esta todos los dias y un `grave` que esta en uno
    tipos = [{"tipo": "ruido", "severidad": "info", "fuente": "f", "dias": 30},
             {"tipo": "rango", "severidad": "grave", "fuente": "f", "dias": 1}]

    # When
    lista = hechos.armar("a", "b", [], [], secciones.fila_agregada("Total", [], None),
                         tipos)

    # Then
    destacados = [h.indicador for h in lista if "hallazgo" in h.indicador]
    assert "rango" in destacados[0]


# ── Verificador ─────────────────────────────────────────────────────────────────
HECHOS = [Hecho("H1", "Días con datos", 57, "días"),
          Hecho("H2", "Energía AC", 363.66, "kWh", "", 2),
          Hecho("H3", "PR de 2026-07", None)]


def test_las_marcas_se_resuelven_con_valor_y_unidad():
    # Given
    parrafos = [{"tipo": "hecho", "texto": "Hubo datos en [[H1]] y se generó [[H2]]."}]

    # When
    revision = verificar.revisar(parrafos, HECHOS, [])

    # Then
    assert revision.aprobada
    assert revision.parrafos[0]["texto"] == "Hubo datos en 57 días y se generó 363,66 kWh."
    assert revision.parrafos[0]["citas"] == ["H1", "H2"]


def test_una_cifra_que_no_esta_en_los_hechos_rechaza_la_lectura():
    # Given: el modelo escribe un numero por su cuenta
    parrafos = [{"tipo": "hecho", "texto": "Hubo datos en [[H1]], un 63 % del periodo."}]

    # When
    revision = verificar.revisar(parrafos, HECHOS, [])

    # Then
    assert not revision.aprobada
    assert "63" in revision.problemas[0]


def test_una_cita_a_un_hecho_que_no_existe_rechaza_la_lectura():
    revision = verificar.revisar([{"tipo": "hecho", "texto": "Fueron [[H99]]."}],
                                 HECHOS, [])
    assert not revision.aprobada and "H99" in revision.problemas[0]


def test_un_hecho_sin_ninguna_cita_se_rechaza():
    revision = verificar.revisar([{"tipo": "hecho", "texto": "La planta anduvo bien."}],
                                 HECHOS, [])
    assert not revision.aprobada


def test_nombres_con_digitos_y_cifras_de_las_advertencias_se_aceptan():
    # Given: PV1 es un nombre; 2026 y 07 estan en un indicador; 722 en una advertencia
    parrafos = [
        {"tipo": "hecho", "texto": "En PV1 hubo [[H1]]; el PR de 2026-07 es [[H3]]."},
        {"tipo": "hipotesis", "texto": "Hipótesis: calibrar contra el SP722 lo aclara."},
    ]

    # When
    revision = verificar.revisar(parrafos, HECHOS, ["Falta calibrar con el SP722."])

    # Then
    assert revision.aprobada
    assert "sin dato" in revision.parrafos[0]["texto"]


def test_la_unidad_repetida_por_el_modelo_se_limpia():
    revision = verificar.revisar(
        [{"tipo": "hecho", "texto": "Hubo datos en [[H1]] días y [[H2]] kWh."}], HECHOS, [])
    assert revision.parrafos[0]["texto"] == "Hubo datos en 57 días y 363,66 kWh."


# ── Redaccion, con un cliente falso ─────────────────────────────────────────────
class _ClienteFalso:
    """Devuelve una respuesta por llamada y guarda lo que le pidieron."""

    def __init__(self, *lecturas, stop_reason="end_turn"):
        self.pendientes = list(lecturas)
        self.pedidos = []
        self.stop_reason = stop_reason
        self.messages = self

    def create(self, **kwargs):
        self.pedidos.append(kwargs)
        cuerpo = self.pendientes.pop(0)
        if isinstance(cuerpo, Exception):
            raise cuerpo
        return SimpleNamespace(
            content=[SimpleNamespace(type="text", text=json.dumps({"parrafos": cuerpo}))],
            usage=SimpleNamespace(input_tokens=100, output_tokens=50),
            stop_reason=self.stop_reason)


BUENA = [{"tipo": "hecho", "texto": "Hubo datos en [[H1]]."}]
MALA = [{"tipo": "hecho", "texto": "Hubo datos en 99 días."}]


def test_una_lectura_valida_se_publica_al_primer_intento():
    # Given
    cliente = _ClienteFalso(BUENA)

    # When
    lectura = redaccion.redactar(HECHOS, [], client=cliente)

    # Then: el modelo no recibio tools ni `tool_choice`, solo el esquema de salida
    assert lectura["intentos"] == 1 and lectura["parrafos"][0]["texto"] == "Hubo datos en 57 días."
    pedido = cliente.pedidos[0]
    assert "tool_choice" not in pedido and pedido["output_config"]["format"]["type"] == "json_schema"
    assert lectura["usage"] == {"input_tokens": 100, "output_tokens": 50, "requests": 1}


def test_una_lectura_rechazada_se_pide_de_nuevo_con_los_problemas():
    # Given
    cliente = _ClienteFalso(MALA, BUENA)

    # When
    lectura = redaccion.redactar(HECHOS, [], client=cliente)

    # Then: el segundo pedido lleva el turno del modelo y el rechazo con la cifra
    assert lectura["intentos"] == 2 and lectura["parrafos"]
    mensajes = cliente.pedidos[1]["messages"]
    assert mensajes[1]["role"] == "assistant" and "99" in mensajes[2]["content"]


def test_dos_lecturas_rechazadas_dejan_el_informe_sin_lectura_y_con_motivo():
    lectura = redaccion.redactar(HECHOS, [], client=_ClienteFalso(MALA, MALA))
    assert lectura["parrafos"] == [] and "99" in lectura["motivo"]
    assert lectura["intentos"] == redaccion.INTENTOS


def test_un_fallo_del_modelo_no_levanta():
    lectura = redaccion.redactar(HECHOS, [], client=_ClienteFalso(RuntimeError("sin red")))
    assert lectura["parrafos"] == [] and "sin red" in lectura["motivo"]


def test_un_rechazo_del_modelo_queda_como_motivo():
    lectura = redaccion.redactar(HECHOS, [],
                                 client=_ClienteFalso(BUENA, stop_reason="refusal"))
    assert lectura["parrafos"] == [] and "declinó" in lectura["motivo"]


def test_el_foco_viaja_como_tema_y_no_reemplaza_los_hechos():
    cliente = _ClienteFalso(BUENA)
    redaccion.redactar(HECHOS, ["aviso"], foco="el inversor", client=cliente)
    pedido = cliente.pedidos[0]["messages"][0]["content"]
    assert "el inversor" in pedido and '"H1"' in pedido and "- aviso" in pedido


# ── Libro y punta a punta ───────────────────────────────────────────────────────
@pytest.fixture
def sin_base(monkeypatch, insumos):
    monkeypatch.setattr(flujo.datos, "recolectar", lambda v: insumos)
    monkeypatch.setattr(flujo.resumen, "ultimo_global", lambda: "2026-07-02T17:00:00")


def _libro(cuerpo: bytes):
    return openpyxl.load_workbook(io.BytesIO(cuerpo))


def test_el_libro_trae_las_hojas_y_la_cabecera_en_la_fila_uno(sin_base):
    # Given/When
    resultado_ = informe.generar("2026-07-01", "2026-07-03", client=_ClienteFalso(BUENA))
    wb = _libro(resultado_.cuerpo)

    # Then
    assert wb.sheetnames == ["Resumen", "Diario", "Mensual", "Disponibilidad",
                             "Calidad", "Gráficos", "Método"]
    diario = list(wb["Diario"].iter_rows(values_only=True))
    assert diario[0][:4] == ("fecha", "filas_electrico", "filas_radiacion",
                             "energia_ac_kwh")
    assert len(diario) == 4                      # cabecera + tres dias
    assert diario[3][3] is None                  # el dia sin datos: vacio, no cero
    assert resultado_.nombre == "informe-agrovoltaic-sc_2026-07-01_2026-07-03.xlsx"


def test_el_rango_incluye_el_ultimo_dia():
    # Given/When
    v, desde, hasta = flujo.ventana_inclusiva("2026-07-01", "2026-07-03")

    # Then: `hasta` exclusivo de la ventana es el dia SIGUIENTE
    assert v.sql == ("2026-07-01", "2026-07-04") and (desde, hasta) == ("2026-07-01", "2026-07-03")


def test_la_lectura_y_sus_citas_quedan_en_el_resumen(sin_base):
    wb = _libro(informe.generar("2026-07-01", "2026-07-03",
                                client=_ClienteFalso([{"tipo": "hecho",
                                                       "texto": "Periodo desde [[H1]]."}])).cuerpo)
    celdas = [c for fila in wb["Resumen"].iter_rows(values_only=True) for c in fila if c]
    assert "Periodo desde 2026-07-01." in celdas and "H1" in celdas


def test_sin_lectura_el_resumen_dice_por_que(sin_base):
    resultado_ = informe.generar("2026-07-01", "2026-07-03",
                                 client=_ClienteFalso(RuntimeError("sin red")))
    celdas = [c for fila in _libro(resultado_.cuerpo)["Resumen"].iter_rows(values_only=True)
              for c in fila if c]
    assert any(str(c).startswith("Sin lectura:") and "sin red" in str(c) for c in celdas)


def test_el_metodo_define_cada_columna_de_cada_hoja(sin_base, composicion):
    wb = _libro(informe.generar("2026-07-01", "2026-07-03", con_lectura=False).cuerpo)
    definidas = {(f[0], f[2]) for f in wb["Método"].iter_rows(values_only=True) if f[2]}
    for hoja in composicion["hojas"]:
        for columna in hoja.columnas:
            assert (hoja.nombre, columna.clave) in definidas


# ── Endpoint y tool ─────────────────────────────────────────────────────────────
def test_el_endpoint_devuelve_un_xlsx_adjunto_y_dice_si_trae_lectura(sin_base):
    # Given/When: sin lectura no se toca el modelo
    r = CLIENTE.get("/informe", params={"desde": "2026-07-01", "hasta": "2026-07-03",
                                        "lectura": "false"})

    # Then
    assert r.status_code == 200
    assert r.headers["content-type"] == informe.libro.MIME
    assert "informe-agrovoltaic-sc_2026-07-01_2026-07-03.xlsx" in r.headers["content-disposition"]
    assert r.headers[api.CABECERA_LECTURA] == api.SIN_LECTURA
    assert _libro(r.content).sheetnames[0] == "Resumen"


def test_el_endpoint_exige_las_dos_fechas():
    r = CLIENTE.get("/informe", params={"desde": "2026-07-01"})
    assert r.status_code in (400, 422)


def test_el_endpoint_con_lectura_pasa_por_el_freno_de_consumo(sin_base, monkeypatch):
    # Given: presupuesto agotado
    monkeypatch.setattr(api.limites, "presupuesto_agotado", lambda: (True, 5.0, 5.0))

    # When/Then: con lectura se frena, sin lectura sigue saliendo
    assert CLIENTE.get("/informe", params={"desde": "2026-07-01",
                                           "hasta": "2026-07-03"}).status_code == 429
    assert CLIENTE.get("/informe", params={"desde": "2026-07-01", "hasta": "2026-07-03",
                                           "lectura": "false"}).status_code == 200


def test_la_tool_devuelve_los_hechos_y_la_ruta_de_descarga(sin_base):
    # Given/When
    salida = preparar_informe.run("2026-07-01", "2026-07-03", foco="el inversor")

    # Then
    assert salida["hechos"][0] == {"id": "H1", "indicador": "Inicio del periodo",
                                   "valor": "2026-07-01"}
    assert salida["descarga"]["ruta"] == (
        "/informe?desde=2026-07-01&hasta=2026-07-03&foco=el+inversor")
    assert salida["periodo"]["hasta_inclusivo"] is True
