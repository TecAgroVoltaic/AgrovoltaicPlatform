"""Estilos y bloques de celdas que comparten las hojas del libro. Sin logica."""
from __future__ import annotations

from openpyxl.styles import Alignment, Font, PatternFill

NEGRITA = Font(bold=True)
TITULO = Font(bold=True, size=14)
CABECERA = PatternFill("solid", fgColor="DDE5ED")
AJUSTE = Alignment(wrap_text=True, vertical="top")
ANCHO_MIN, ANCHO_MAX = 10, 44
ANCHO_TEXTO = 110
ALTO_GRAFICO, ANCHO_GRAFICO = 8.5, 26.0
FILAS_POR_GRAFICO = 19

TOTAL = "Total"


def cabecera(ws, fila: int, titulos: list[str]) -> None:
    for i, titulo in enumerate(titulos, start=1):
        c = ws.cell(row=fila, column=i, value=titulo)
        c.font, c.fill = NEGRITA, CABECERA


def seccion(ws, fila: int, titulo: str) -> int:
    ws.cell(row=fila, column=1, value=titulo).font = NEGRITA
    return fila + 1
