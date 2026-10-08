"""El libro de Excel. Solo escribe: no calcula ni decide nada.

Las hojas de datos llevan la CABECERA EN LA FILA 1 y el nombre de columna en
snake_case con la unidad en el nombre (`energia_ac_kwh`). Es a proposito: asi
`readtable` de MATLAB y `pandas.read_excel` las leen sin saltar filas ni renombrar.
La etiqueta legible, la unidad y la definicion de cada columna estan en la hoja
Método, y los graficos en una hoja aparte para no ensuciar los datos.

  libro_estilo.py    fuentes, anchos y bloques de celdas compartidos.
  libro_texto.py     las hojas Resumen y Método.
  libro_graficos.py  la hoja Gráficos.
"""
from __future__ import annotations

import io
from datetime import date

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from historico.informe.libro_estilo import (ANCHO_MAX, ANCHO_MIN, NEGRITA, TOTAL,
                                            cabecera)
from historico.informe.libro_graficos import graficos
from historico.informe.libro_texto import metodo, resumen
from historico.informe.tabla import FECHA, Hecho, Hoja

MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

__all__ = ["MIME", "TOTAL", "escribir"]


def _celda(valor, formato: str | None):
    if formato == FECHA and isinstance(valor, str) and valor:
        return date.fromisoformat(valor[:10])
    return valor


def _hoja_de_datos(wb: Workbook, hoja: Hoja) -> None:
    ws = wb.create_sheet(hoja.nombre)
    cabecera(ws, 1, [c.clave for c in hoja.columnas])
    for f, fila in enumerate(hoja.filas, start=2):
        for i, col in enumerate(hoja.columnas, start=1):
            c = ws.cell(row=f, column=i, value=_celda(fila.get(col.clave), col.formato))
            if col.formato:
                c.number_format = col.formato
            if fila.get(hoja.columnas[0].clave) == TOTAL:
                c.font = NEGRITA
    for i, col in enumerate(hoja.columnas, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(
            ANCHO_MIN, min(ANCHO_MAX, len(col.clave) + 2))
    ws.freeze_panes = "B2"
    if hoja.filas:
        ws.auto_filter.ref = ws.dimensions


def escribir(meta: dict, lectura: dict, advertencias: list[str], hechos: list[Hecho],
             hojas: list[Hoja], parametros: list[tuple[str, str]],
             traza: list[tuple[str, str]]) -> bytes:
    """El .xlsx completo, en memoria. `hojas` va en el orden en que se leen."""
    wb = Workbook()
    resumen(wb, meta, lectura, advertencias, hechos)
    for hoja in hojas:
        _hoja_de_datos(wb, hoja)
    por_nombre = {h.nombre: h for h in hojas}
    graficos(wb, por_nombre["Diario"], por_nombre["Mensual"])
    metodo(wb, hojas, parametros, traza)
    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()
